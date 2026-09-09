#!/usr/bin/env python3
"""Render docs/thesis_draft.md into an editable Word document.

Re-run this after editing the markdown; the .docx is a build artefact, not the
source. Figures named by `*[Insert Figure N - path]*` placeholders are embedded
from disk, so regenerating the diagrams and re-running this picks them up.

  python scripts/build_thesis_docx.py
  python scripts/build_thesis_docx.py --input docs/thesis_draft.md --output out.docx
"""
import argparse
import io
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
BODY_SIZE = Pt(12)
LINE_SPACING = 1.5

# Page area available for a figure, in inches. Tall diagrams are capped by
# height rather than width so they cannot run off the bottom of the page.
MAX_FIG_W = 6.0
MAX_FIG_H = 8.0

FIG_PLACEHOLDER = re.compile(r"^\*\[Insert (?:Figure|Table)\s+([\d.]+)\s*[-—]\s*`?([^`\]]+?)`?\]\*$")
CAPTION = re.compile(r"^\*((?:Figure|Table)\s+[\d.]+[:.].*)\*$", re.S)
INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`)")


def set_style(doc):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    # Ensure the East Asian font mapping also points at the body font, or Word
    # silently substitutes for some characters.
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_after = Pt(6)

    for name, size, bold in (
        ("Heading 1", 16, True),
        ("Heading 2", 14, True),
        ("Heading 3", 12, True),
    ):
        st = doc.styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.0


def add_runs(par, text):
    """Add text to a paragraph, honouring **bold**, *italic* and `code`."""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            par.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            r = par.add_run(chunk[1:-1])
            r.font.name = MONO_FONT
            r.font.size = Pt(10.5)
        elif chunk.startswith("*") and chunk.endswith("*"):
            par.add_run(chunk[1:-1]).italic = True
        else:
            par.add_run(chunk)


def emit_table(doc, rows):
    """Render a markdown table block as a Word table."""
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    # Drop the |---|---| separator row.
    body = [c for c in cells if not all(set(x) <= set("-: ") and x for x in c)]
    if not body:
        return
    width = max(len(r) for r in body)
    table = doc.add_table(rows=len(body), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(body):
        for j in range(width):
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.space_after = Pt(2)
            text = row[j] if j < len(row) else ""
            add_runs(par, text)
            for run in par.runs:
                run.font.size = Pt(10)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def emit_image(doc, path, fig_id):
    full = path if os.path.isabs(path) else os.path.join(ROOT, path)
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not os.path.exists(full):
        r = par.add_run(f"[Figure {fig_id} missing: {path}]")
        r.italic = True
        print(f"  WARNING: missing image {path}")
        return
    from PIL import Image
    with Image.open(full) as im:
        w, h = im.size
    scale = min(MAX_FIG_W / (w / 96.0), MAX_FIG_H / (h / 96.0), 1.0
                if w / 96.0 <= MAX_FIG_W and h / 96.0 <= MAX_FIG_H else 99)
    disp_w = min(MAX_FIG_W, (w / 96.0) * scale)
    par.add_run().add_picture(full, width=Inches(disp_w))


def build(md_path, out_path):
    lines = io.open(md_path, encoding="utf-8").read().split("\n")
    doc = Document()
    set_style(doc)
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.25)
        s.right_margin = Inches(1.0)

    i = 0
    first_heading = True
    n_img = n_tab = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip() or line.strip() == "---":
            i += 1
            continue

        # Tables
        if line.lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            emit_table(doc, block)
            n_tab += 1
            continue

        # Fenced code
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            par = doc.add_paragraph()
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.left_indent = Inches(0.4)
            r = par.add_run("\n".join(code))
            r.font.name = MONO_FONT
            r.font.size = Pt(10)
            continue

        # Headings
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            if level == 1 and not first_heading:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first_heading = False
            doc.add_heading(re.sub(r"[*`]", "", text), level=level)
            i += 1
            continue

        # Figure placeholder
        m = FIG_PLACEHOLDER.match(line.strip())
        if m:
            emit_image(doc, m.group(2).strip(), m.group(1))
            n_img += 1
            i += 1
            continue

        # Caption (may wrap over several lines)
        if line.strip().startswith("*Figure") or line.strip().startswith("*Table"):
            buf = [line.strip()]
            while not buf[-1].endswith("*") or len(buf[-1]) < 2:
                i += 1
                if i >= len(lines):
                    break
                buf.append(lines[i].strip())
            text = " ".join(buf)
            cm = CAPTION.match(text)
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.0
            r = par.add_run(cm.group(1) if cm else text.strip("*"))
            r.italic = True
            r.font.size = Pt(10.5)
            i += 1
            continue

        # Lists
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            style = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            body = [m.group(3)]
            i += 1
            while i < len(lines) and lines[i].startswith("   ") and lines[i].strip() \
                    and not re.match(r"^\s*([-*]|\d+\.)\s", lines[i]):
                body.append(lines[i].strip())
                i += 1
            par = doc.add_paragraph(style=style)
            if indent >= 2:
                par.paragraph_format.left_indent = Inches(0.6)
            add_runs(par, " ".join(body))
            continue

        # Paragraph — join wrapped lines
        body = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#{1,3}\s|\||```|\s*([-*]|\d+\.)\s|\*Figure|\*Table|\*\[Insert)", lines[i]
        ) and lines[i].strip() != "---":
            body.append(lines[i].strip())
            i += 1
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(par, " ".join(body))

    doc.save(out_path)
    print(f"Wrote {out_path}")
    print(f"  {n_img} figures embedded, {n_tab} tables rendered")


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--input", default=os.path.join(ROOT, "docs", "thesis_draft.md"))
    p.add_argument("--output", default=os.path.join(ROOT, "docs", "thesis_draft.docx"))
    a = p.parse_args()
    build(a.input, a.output)


if __name__ == "__main__":
    main()
